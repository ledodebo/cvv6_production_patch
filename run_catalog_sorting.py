import json
import os
import re
import shutil
import time
import openpyxl

INSTRUCTIONS_PATH = r'E:\xampp\htdocs\CVV6\catalog_sorting_instructions.json'
EXCEL_PATH = r'E:\xampp\htdocs\CVV6\foxreload_direct_catalog.xlsx'
BACKUP_PATH = r'E:\xampp\htdocs\CVV6\foxreload_direct_catalog.backup.xlsx'

print("=== GBC Catalog AI Sorting Engine ===")
start_time = time.time()

# 1. Load instructions JSON
with open(INSTRUCTIONS_PATH, 'r', encoding='utf-8') as f:
    instructions = json.load(f)

print(f"Loaded instructions: {instructions.get('title')} (v{instructions.get('version')})")

# Extract curated taxonomies
master_brands_config = instructions.get('curated_taxonomies', {}).get('multi_country_master_brands', [])
MASTER_BRANDS = {}
for mb in master_brands_config:
    b_key = mb['brand'].lower().strip()
    MASTER_BRANDS[b_key] = mb['clean_product_name']

service_config = instructions.get('curated_taxonomies', {}).get('service_subscription_plan_splitting_rules', {})
SERVICE_BRANDS = set(b.lower().strip() for b in service_config.get('brands', []))

distinct_editions = instructions.get('curated_taxonomies', {}).get('distinct_edition_brands', [])

# Additional known multi-country brands from our curated domain mapping
ADDITIONAL_MASTER_BRANDS = {
    'pubg mobile': 'PUBG Mobile',
    'free fire': 'Free Fire',
    'mobile legends': 'Mobile Legends',
    'bilibili': 'Bilibili',
    'tinder': 'Tinder',
    'valorant': 'Valorant',
    'apex legends': 'Apex Legends',
    'brawl stars': 'Brawl Stars',
    'clash of clans': 'Clash of Clans',
    'clash royale': 'Clash Royale',
    'call of duty': 'Call of Duty',
    'arena of valor': 'Arena of Valor',
    'honor of kings': 'Honor of Kings',
    'bigo': 'Bigo Live',
    'bigo live': 'Bigo Live',
    'garena': 'Garena Shells',
    'garena shells': 'Garena Shells',
    'roblox': 'Roblox',
    'league of legends': 'League of Legends',
    'playstation': 'PlayStation Network',
    'xbox': 'Xbox Live',
    'nintendo': 'Nintendo eShop',
    'apple': 'Apple Gift Card',
    'google play': 'Google Play',
    'steam': 'Steam Wallet',
    'razer gold': 'Razer Gold',
}
for k, v in ADDITIONAL_MASTER_BRANDS.items():
    if k not in MASTER_BRANDS:
        MASTER_BRANDS[k] = v

# Known single-region services where plans must split
ADDITIONAL_SERVICE_BRANDS = {
    'suno ai', 'perplexity', 'claude', 'clipdrop', 'd-id', 'dzine.ai',
    'elevenlabs', 'fal.ai', 'flora ai', 'gamma.app', 'genspark', 'grok',
    'hedra', 'heyboss.ai', 'higgsfield ai', 'ideogram', 'jetbrains ai',
    'kimi ai', 'kling', 'krea', 'lovable', 'lovart.ai', 'magnific',
    'midjourney', 'otter.ai', 'chatgpt', 'discord', 'canva', 'cursor'
}
SERVICE_BRANDS.update(ADDITIONAL_SERVICE_BRANDS)

# Plan tier patterns (Law 2)
TIER_PATTERNS = [
    ('Nitro Basic', r'\bnitro basic\b'),
    ('Nitro', r'\bnitro\b'),
    ('Server Boost', r'\bserver boost\b'),
    ('Premier', r'\bpremier\b'),
    ('Plus', r'\bplus\b'),
    ('Pro Plan', r'\bpro plan\b'),
    ('Basic Plan', r'\bbasic plan\b'),
    ('Standard Plan', r'\bstandard plan\b'),
    ('Mega Plan', r'\bmega plan\b'),
    ('Max 20x', r'\bmax 20x\b'),
    ('Pro', r'\bpro\b'),
    ('Ultra', r'\bultra\b'),
    ('Super', r'\bsuper\b'),
    ('Business', r'\bbusiness\b'),
    ('Creator', r'\bcreator\b'),
    ('Starter', r'\bstarter\b'),
    ('Basic', r'\bbasic\b'),
    ('Lite', r'\blite\b'),
    ('Beginner', r'\bbeginner\b'),
    ('Master', r'\bmaster\b'),
    ('Ultimate', r'\bultimate\b'),
    ('Premium+', r'\bpremium\+\b'),
    ('Premium', r'\bpremium\b'),
]

def sanitize_title(title):
    """Applies Law 4: Title Sanitization"""
    if not title:
        return ""
    t = str(title).strip()
    # Strip prefixes like Payment cards for services (X)
    t = re.sub(r'^(?:Payment cards for services|Activation codes|Gift Cards|Digital Codes|Direct Top-Up|Top-Up)\s*\(\s*', '', t, flags=re.IGNORECASE)
    t = re.sub(r'\s*\)\s*$', '', t)
    # Strip tags
    t = re.sub(r'#(?:2|3|4|5)\b', '', t)
    t = re.sub(r'\[Instant Delivery\]', '', t, flags=re.IGNORECASE)
    t = re.sub(r'\((?:Key|Code|Global)\)', '', t, flags=re.IGNORECASE)
    t = re.sub(r'\s+', ' ', t).strip(" -:_()")
    return t

def split_service_item(brand_name, item_name):
    """Applies Law 2: Single-Region Plan Separation"""
    raw = item_name or ""
    has_vat = bool(re.search(r'\(incl\.?\s*27%\s*vat\)', raw, re.IGNORECASE))
    clean_no_vat = re.sub(r'\(incl\.?\s*27%\s*vat\)', '', raw, flags=re.IGNORECASE).strip()
    clean_name = re.sub(re.escape(brand_name), '', clean_no_vat, flags=re.IGNORECASE).strip(" -:_()")
    is_renewal = bool(re.search(r'\brenewal\b', raw, re.IGNORECASE))
    clean_no_renewal = re.sub(r'\brenewal\b', '', clean_name, flags=re.IGNORECASE).strip(" -:_()")

    matched_tier = None
    for tier_name, pattern in TIER_PATTERNS:
        if re.search(pattern, clean_no_renewal, re.IGNORECASE):
            matched_tier = tier_name
            break

    if matched_tier:
        plan_label = f"Renewal {matched_tier}" if is_renewal else matched_tier
        prod_title = f"{brand_name} - {plan_label}"
    else:
        prod_title = brand_name

    dur_match = re.search(r'(\d+\s*(?:month|year|day|credit)s?)', clean_no_vat, re.IGNORECASE)
    duration = dur_match.group(1).title() if dur_match else ""

    if has_vat:
        var_label = f"{duration} (With 27% VAT)".strip() if duration else "With 27% VAT"
    elif duration:
        var_label = duration
    else:
        var_label = "Standard"

    return prod_title, var_label

# 2. Backup Excel file
print(f"Creating backup at {BACKUP_PATH}...")
shutil.copy2(EXCEL_PATH, BACKUP_PATH)

# 3. Load Excel workbook
print(f"Loading {EXCEL_PATH}...")
wb = openpyxl.load_workbook(EXCEL_PATH)
ws = wb.active

header = [cell.value for cell in ws[1]]
h_map = {str(name).strip(): idx for idx, name in enumerate(header) if name}

cat_col = h_map['Category'] + 1
brand_col = h_map['Brand'] + 1
pname_col = h_map['Product Name'] + 1
master_col = h_map['Is Master Product'] + 1
cname_col = h_map['Country Name'] + 1
ccode_col = h_map['Country Code'] + 1
vname_col = h_map['Variation Name'] + 1

total_rows = ws.max_row - 1
print(f"Total catalog variation rows: {total_rows}")

master_updated_count = 0
service_split_count = 0
sanitized_count = 0

# Track master products and split service products
master_products_seen = set()
split_services_seen = set()

for row in range(2, ws.max_row + 1):
    brand = str(ws.cell(row=row, column=brand_col).value or '').strip()
    brand_lower = brand.lower()
    p_name = str(ws.cell(row=row, column=pname_col).value or '').strip()
    v_name = str(ws.cell(row=row, column=vname_col).value or '').strip()
    c_name = str(ws.cell(row=row, column=cname_col).value or '').strip()
    is_master = str(ws.cell(row=row, column=master_col).value or '').strip()

    # Rule 4: Sanitize title wrapper
    clean_p_name = sanitize_title(p_name)
    if clean_p_name != p_name:
        ws.cell(row=row, column=pname_col, value=clean_p_name)
        p_name = clean_p_name
        sanitized_count += 1

    # Rule 1: Multi-Country Master Consolidation
    if brand_lower in MASTER_BRANDS:
        canonical_master_name = MASTER_BRANDS[brand_lower]
        if p_name != canonical_master_name or is_master != 'Yes':
            ws.cell(row=row, column=pname_col, value=canonical_master_name)
            ws.cell(row=row, column=master_col, value='Yes')
            master_updated_count += 1
        master_products_seen.add(canonical_master_name)

    # Rule 2: Single-Region Plan Separation for Services
    elif brand_lower in SERVICE_BRANDS or (c_name.lower() in ['global', 'gl', ''] and any(sb in brand_lower for sb in SERVICE_BRANDS)):
        item_raw = f"{brand} {v_name}" if brand_lower not in v_name.lower() else v_name
        new_prod_title, new_var_label = split_service_item(brand, item_raw)

        if p_name != new_prod_title or v_name != new_var_label or is_master != 'No':
            ws.cell(row=row, column=pname_col, value=new_prod_title)
            ws.cell(row=row, column=vname_col, value=new_var_label)
            ws.cell(row=row, column=master_col, value='No')
            service_split_count += 1
        split_services_seen.add(new_prod_title)

# 4. Save workbook
print("Saving normalized catalog to Excel...")
wb.save(EXCEL_PATH)
wb.close()

# Also update the production patch copy if it exists
patch_excel = r'E:\xampp\htdocs\CVV6\cvv6_production_patch\foxreload_direct_catalog.xlsx'
try:
    shutil.copy2(EXCEL_PATH, patch_excel)
except Exception:
    pass

elapsed = round(time.time() - start_time, 2)
print(f"\nCatalog Sorting Complete in {elapsed} seconds!")
print(f"--------------------------------------------------")
print(f"Total Rows Checked: {total_rows}")
print(f"Master Product Rows Consolidated (Law 1): {master_updated_count}")
print(f"  Distinct Master Brands: {len(master_products_seen)} (e.g. Bilibili, Tinder, PUBG, Free Fire, Valorant)")
print(f"Single-Region Service Rows Split (Law 2): {service_split_count}")
print(f"  Distinct Product Cards Created: {len(split_services_seen)} (e.g. Suno AI - Premier, Midjourney - Pro Plan)")
print(f"Titles Sanitized (Law 4): {sanitized_count}")
print(f"Saved: {EXCEL_PATH}")
print(f"Backup: {BACKUP_PATH}")
